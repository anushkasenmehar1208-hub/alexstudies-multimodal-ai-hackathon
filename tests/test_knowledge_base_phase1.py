"""Phase 1 ingestion tests against the isolated local PostgreSQL database."""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import sys
import unittest

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import create_engine, delete, inspect
from sqlalchemy.engine import make_url

ROOT = Path(__file__).parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from uni_app import youtube_utils  # noqa: E402
from tools.pdf_fixture import build_pdf  # noqa: E402
from uni_app.knowledge_base import (  # noqa: E402
    DeferredVisualAnalyzer,
    DuplicateSourceError,
    InvalidSourceError,
    KnowledgeBaseIngestor,
    SourceAsset,
    SourceBase,
)


OWNER_A = 1_901_001
OWNER_B = 1_901_002


def _local_engine():
    raw_url = os.environ.get("DATABASE_URL", "")
    if not raw_url:
        raise RuntimeError("DATABASE_URL must point to the isolated Phase 1 database")
    url = make_url(raw_url)
    if url.host not in {"127.0.0.1", "localhost"} or url.port != 55432:
        raise RuntimeError("refusing to run knowledge-base tests outside local port 55432")
    if url.database != "alexstudies_hackathon":
        raise RuntimeError("refusing to run knowledge-base tests outside alexstudies_hackathon")
    return create_engine(url.set(drivername="postgresql+psycopg"), pool_pre_ping=True)


def _png_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (18, 12), color=(20, 120, 210)).save(output, format="PNG")
    return output.getvalue()


def _pdf_fixture() -> bytes:
    return build_pdf(
        [
            (320, 420, "Page one: energy is conserved.", None),
            (500, 360, "Page two: this page contains a figure.", _png_bytes()),
        ]
    )


def _pptx_fixture() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])
    first.shapes.title.text = "Forces and motion"
    table = first.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Quantity"
    table.cell(0, 1).text = "Unit"
    table.cell(1, 0).text = "Force"
    table.cell(1, 1).text = "Newton"
    first.notes_slide.notes_text_frame.text = "Explain the free-body diagram before the example."

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Worked example"
    second.shapes.add_picture(BytesIO(_png_bytes()), Inches(1), Inches(1.5), width=Inches(1.5))
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class KnowledgeBasePhase1Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = _local_engine()
        SourceBase.metadata.create_all(cls.engine)
        cls.ingestor = KnowledgeBaseIngestor(cls.engine)

    @classmethod
    def tearDownClass(cls):
        with cls.engine.begin() as connection:
            connection.execute(delete(SourceAsset).where(SourceAsset.owner_id.in_([OWNER_A, OWNER_B])))
        cls.engine.dispose()

    def setUp(self):
        with self.engine.begin() as connection:
            connection.execute(delete(SourceAsset).where(SourceAsset.owner_id.in_([OWNER_A, OWNER_B])))

    def test_schema_contains_all_phase_one_models(self):
        tables = set(inspect(self.engine).get_table_names())
        self.assertTrue(
            {
                "source_assets",
                "source_units",
                "source_visuals",
                "topics",
                "topic_prerequisites",
                "source_unit_topics",
            }.issubset(tables)
        )

    def test_youtube_segments_keep_exact_timestamp_ranges(self):
        segments = [
            youtube_utils.TranscriptSegment("Energy cannot be created.", 0.125, 2.75),
            youtube_utils.TranscriptSegment("It changes form.", 2.75, 4.5),
        ]
        asset = self.ingestor.ingest_youtube(
            owner_id=OWNER_A,
            original_url="https://www.youtube.com/watch?v=abcdefghijk",
            title="Energy lecture",
            segments=segments,
        )
        units = self.ingestor.list_units(owner_id=OWNER_A, asset_id=asset.id)
        self.assertEqual([unit.extracted_text for unit in units], [segment.text for segment in segments])
        self.assertEqual(units[0].source_locator["start_seconds"], 0.125)
        self.assertEqual(units[0].source_locator["end_seconds"], 2.75)
        self.assertEqual(units[1].source_locator["start_seconds"], 2.75)
        self.assertEqual(units[1].source_locator["end_seconds"], 4.5)
        self.assertEqual(asset.original_url, "https://www.youtube.com/watch?v=abcdefghijk")
        self.assertEqual(asset.external_id, "abcdefghijk")

    def test_supadata_normalizer_preserves_millisecond_offsets(self):
        segments = youtube_utils._segments_from_supadata_content(
            [{"text": "First", "offset": 1250, "duration": 750}]
        )
        self.assertEqual(segments, [youtube_utils.TranscriptSegment("First", 1.25, 2.0)])

    def test_pdf_keeps_pages_dimensions_and_embedded_figure_relationship(self):
        asset = self.ingestor.ingest_pdf(
            owner_id=OWNER_A,
            filename="physics.pdf",
            data=_pdf_fixture(),
        )
        units = self.ingestor.list_units(owner_id=OWNER_A, asset_id=asset.id)
        visuals = self.ingestor.list_visuals(owner_id=OWNER_A, asset_id=asset.id)
        self.assertEqual([unit.source_locator["page_number"] for unit in units], [1, 2])
        self.assertEqual(units[0].source_locator["width_points"], 320.0)
        self.assertEqual(units[1].source_locator["height_points"], 360.0)
        self.assertIn("energy is conserved", units[0].extracted_text)
        self.assertEqual(len(visuals), 1)
        self.assertEqual(visuals[0].source_locator["page_number"], 2)
        self.assertEqual(visuals[0].source_unit_id, units[1].id)
        self.assertGreater(len(visuals[0].image_bytes), 20)
        self.assertEqual(visuals[0].analysis_status, "not_requested")
        self.assertIsNone(visuals[0].visual_description)

    def test_pptx_keeps_slides_tables_notes_and_embedded_image(self):
        asset = self.ingestor.ingest_pptx(
            owner_id=OWNER_A,
            filename="mechanics.pptx",
            data=_pptx_fixture(),
        )
        units = self.ingestor.list_units(owner_id=OWNER_A, asset_id=asset.id)
        visuals = self.ingestor.list_visuals(owner_id=OWNER_A, asset_id=asset.id)
        self.assertEqual([unit.source_locator["slide_number"] for unit in units], [1, 2])
        self.assertIn("Forces and motion", units[0].extracted_text)
        self.assertIn("Quantity | Unit", units[0].extracted_text)
        self.assertIn("free-body diagram", units[0].extracted_text)
        self.assertEqual(len(visuals), 1)
        self.assertEqual(visuals[0].source_locator["slide_number"], 2)
        self.assertEqual(visuals[0].source_unit_id, units[1].id)
        self.assertEqual(visuals[0].mime_type, "image/png")

    def test_owner_isolation_and_duplicate_detection(self):
        data = _pdf_fixture()
        first = self.ingestor.ingest_pdf(owner_id=OWNER_A, filename="same.pdf", data=data)
        with self.assertRaises(DuplicateSourceError):
            self.ingestor.ingest_pdf(owner_id=OWNER_A, filename="renamed.pdf", data=data)

        second = self.ingestor.ingest_pdf(owner_id=OWNER_B, filename="same.pdf", data=data)
        self.assertNotEqual(first.id, second.id)
        self.assertIsNone(self.ingestor.get_asset(owner_id=OWNER_B, asset_id=first.id))
        self.assertEqual(self.ingestor.list_units(owner_id=OWNER_B, asset_id=first.id), [])
        self.assertEqual([asset.id for asset in self.ingestor.list_assets(owner_id=OWNER_A)], [first.id])
        self.assertEqual([asset.id for asset in self.ingestor.list_assets(owner_id=OWNER_B)], [second.id])

    def test_invalid_files_and_invalid_video_are_rejected_without_rows(self):
        invalid_calls = (
            lambda: self.ingestor.ingest_pdf(owner_id=OWNER_A, filename="bad.pdf", data=b"not-pdf"),
            lambda: self.ingestor.ingest_pptx(owner_id=OWNER_A, filename="bad.pptx", data=b"not-pptx"),
            lambda: self.ingestor.ingest_youtube(
                owner_id=OWNER_A,
                original_url="https://example.com/video",
                segments=[youtube_utils.TranscriptSegment("text", 0, 1)],
            ),
        )
        for call in invalid_calls:
            with self.subTest(call=call), self.assertRaises(InvalidSourceError):
                call()
        self.assertEqual(self.ingestor.list_assets(owner_id=OWNER_A), [])

    def test_visual_analysis_is_explicitly_deferred(self):
        result = DeferredVisualAnalyzer().analyze(
            image_bytes=_png_bytes(),
            mime_type="image/png",
            source_locator={"kind": "pdf_figure", "page_number": 1},
        )
        self.assertEqual(result.status, "not_requested")
        self.assertIsNone(result.ocr_text)
        self.assertIsNone(result.description)


if __name__ == "__main__":
    unittest.main()
