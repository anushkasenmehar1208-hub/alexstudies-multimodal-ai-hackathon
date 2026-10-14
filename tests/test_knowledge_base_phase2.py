"""Track D HTTP, retrieval, grounding and local OCR integration tests."""

from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
import unittest

import httpx
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import create_engine, delete
from sqlalchemy.engine import make_url
from starlette.applications import Starlette
from starlette.testclient import TestClient

ROOT = Path(__file__).parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from uni_app import ai_provider, youtube_utils  # noqa: E402
from tools.pdf_fixture import build_pdf  # noqa: E402
from uni_app.knowledge_base.api import KnowledgeBaseAPI, add_knowledge_base_routes  # noqa: E402
from uni_app.knowledge_base.models import SourceAsset, SourceBase  # noqa: E402
from uni_app.knowledge_base.visual_analysis import (  # noqa: E402
    DeferredVisualAnalyzer,
    LocalOCRAnalyzer,
)


OWNER = 1_902_001
OTHER_OWNER = 1_902_002
CONFIG = ai_provider.ProviderConfig(
    True,
    "groq",
    "test-only-key",
    "https://api.groq.com/openai/v1",
    "openai/gpt-oss-20b",
    "whisper-large-v3-turbo",
)


def _engine():
    url = make_url(os.environ.get("DATABASE_URL", ""))
    if (url.host, url.port, url.database) not in {
        ("127.0.0.1", 55432, "alexstudies_hackathon"),
        ("localhost", 55432, "alexstudies_hackathon"),
    }:
        raise RuntimeError("refusing to run Phase 2 tests outside the isolated local database")
    return create_engine(url.set(drivername="postgresql+psycopg"), pool_pre_ping=True)


def _text_image(text: str = "ENTROPY DIAGRAM") -> bytes:
    image = Image.new("RGB", (900, 190), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    font = ImageFont.truetype(font_path, 64) if Path(font_path).is_file() else ImageFont.load_default()
    draw.text((30, 55), text, fill="black", font=font)
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _pdf() -> bytes:
    return build_pdf(
        [(500, 500, "Entropy measures the number of microscopic configurations in a system.", _text_image())]
    )


def _pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Entropy in thermodynamics"
    slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(7), Inches(1)).text = (
        "Entropy increases for an isolated system during a spontaneous process."
    )
    slide.shapes.add_picture(BytesIO(_text_image("ENTROPY SLIDE")), Inches(1), Inches(2.5), width=Inches(4))
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _transcript(_video_id: str):
    return [
        youtube_utils.TranscriptSegment(
            "Entropy is associated with the number of accessible microstates.", 12.5, 18.25
        )
    ], ""


def _grounded_completion(_config, messages, **_kwargs):
    identifiers = [int(value) for value in re.findall(r'<source-unit id="(\d+)"', messages[-1]["content"])]
    return json.dumps(
        {
            "answer": "Entropy is connected to accessible microscopic configurations.",
            "supported": True,
            "citations": [{"source_unit_id": identifiers[0]}],
        }
    )


class KnowledgeBasePhase2Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = _engine()
        SourceBase.metadata.create_all(cls.engine)

    @classmethod
    def tearDownClass(cls):
        cls._clean()
        cls.engine.dispose()

    @classmethod
    def _clean(cls):
        with cls.engine.begin() as connection:
            connection.execute(delete(SourceAsset).where(SourceAsset.owner_id.in_([OWNER, OTHER_OWNER])))

    def setUp(self):
        self._clean()
        app = Starlette()
        self.service = KnowledgeBaseAPI(
            engine=self.engine,
            auth_resolver=lambda token: {"owner-token": OWNER, "other-token": OTHER_OWNER}.get(token, -1),
            provider_config=CONFIG,
            completion=_grounded_completion,
            transcript_fetcher=_transcript,
            visual_analyzer=DeferredVisualAnalyzer(),
        )
        add_knowledge_base_routes(app, self.service)
        self.client = TestClient(app)
        self.headers = {"X-Auth-Token": "owner-token"}

    def tearDown(self):
        self.client.close()

    def _ingest_all(self):
        self.pdf_bytes = _pdf()
        pdf = self.client.post(
            "/api/track-d/sources/pdf",
            headers=self.headers,
            files={"file": ("thermodynamics.pdf", self.pdf_bytes, "application/pdf")},
        )
        pptx = self.client.post(
            "/api/track-d/sources/pptx",
            headers=self.headers,
            files={
                "file": (
                    "lecture.pptx",
                    _pptx(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        youtube = self.client.post(
            "/api/track-d/sources/youtube",
            headers=self.headers,
            json={"url": "https://www.youtube.com/watch?v=abcdefghijk", "title": "Entropy lecture"},
        )
        for response in (pdf, pptx, youtube):
            self.assertEqual(response.status_code, 201, response.text)
        return pdf.json(), pptx.json(), youtube.json()

    def test_authenticated_http_ingestion_listing_and_duplicate_handling(self):
        pdf, pptx, youtube = self._ingest_all()
        self.assertEqual((pdf["source_type"], pdf["unit_count"], pdf["visual_count"]), ("pdf", 1, 1))
        self.assertEqual((pptx["source_type"], pptx["unit_count"], pptx["visual_count"]), ("pptx", 1, 1))
        self.assertEqual(youtube["source_metadata"]["segment_count"], 1)
        duplicate = self.client.post(
            "/api/track-d/sources/pdf",
            headers=self.headers,
            files={"file": ("copy.pdf", self.pdf_bytes, "application/pdf")},
        )
        self.assertEqual(duplicate.status_code, 409)
        listed = self.client.get("/api/track-d/sources", headers=self.headers)
        self.assertEqual({item["source_type"] for item in listed.json()["sources"]}, {"pdf", "pptx", "youtube"})

    def test_cross_source_postgres_retrieval_preserves_locations(self):
        self._ingest_all()
        response = self.client.get("/api/track-d/search?q=entropy", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual({item["source_type"] for item in results}, {"pdf", "pptx", "youtube"})
        locators = {item["source_type"]: item["locator"] for item in results}
        self.assertEqual(locators["pdf"]["page_number"], 1)
        self.assertEqual(locators["pptx"]["slide_number"], 1)
        self.assertEqual(locators["youtube"]["start_seconds"], 12.5)
        self.assertTrue(all(item["source_unit_id"] > 0 for item in results))

    def test_tutor_returns_only_server_verified_citation_links(self):
        assets = self._ingest_all()
        expectations = {
            "pdf": ("page 1", "?page=1"),
            "pptx": ("slide 1", "?slide=1"),
            "youtube": ("0:12", "youtube.com/watch?"),
        }
        for asset in assets:
            with self.subTest(source_type=asset["source_type"]):
                response = self.client.post(
                    "/api/track-d/tutor",
                    headers=self.headers,
                    json={"question": "What does this say about entropy?", "asset_ids": [asset["id"]]},
                )
                self.assertEqual(response.status_code, 200, response.text)
                result = response.json()
                self.assertTrue(result["supported"])
                self.assertEqual(len(result["citations"]), 1)
                citation = result["citations"][0]
                label_part, link_part = expectations[asset["source_type"]]
                self.assertIn(label_part, citation["label"])
                self.assertIn(link_part, citation["link"])
                detail = self.client.get(citation["link"], headers=self.headers) if citation["link"].startswith("/") else None
                if detail is not None:
                    self.assertEqual(detail.status_code, 200)
                    self.assertEqual(detail.json()["source_unit_id"], citation["source_unit_id"])

    def test_unsupported_questions_refuse_without_calling_provider(self):
        self._ingest_all()
        self.service.tutor.completion = lambda *_args, **_kwargs: self.fail("provider should not be called")
        response = self.client.post(
            "/api/track-d/tutor",
            headers=self.headers,
            json={"question": "How does chlorophyll absorb blue light?"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["supported"])
        self.assertEqual(response.json()["citations"], [])

    def test_invalid_citation_and_provider_failure_are_safe(self):
        self._ingest_all()
        self.service.tutor.completion = lambda *_args, **_kwargs: json.dumps(
            {"answer": "bad", "supported": True, "citations": [{"source_unit_id": 999999999}]}
        )
        invalid = self.client.post(
            "/api/track-d/tutor", headers=self.headers, json={"question": "Explain entropy"}
        )
        self.assertEqual(invalid.status_code, 503)
        self.assertNotIn("999999999", invalid.text)

        self.service.tutor.completion = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            httpx.ConnectError("test provider down")
        )
        failure = self.client.post(
            "/api/track-d/tutor", headers=self.headers, json={"question": "Explain entropy"}
        )
        self.assertEqual(failure.status_code, 504)
        self.assertEqual(failure.json()["error_code"], "timeout")

    def test_authentication_ownership_invalid_files_and_exact_unit_access(self):
        pdf, _, _ = self._ingest_all()
        self.assertEqual(self.client.get("/api/track-d/sources").status_code, 401)
        self.assertEqual(
            self.client.get(
                "/api/track-d/sources", headers={"Authorization": "Bearer other-token"}
            ).json()["sources"],
            [],
        )
        own_search = self.client.get("/api/track-d/search?q=entropy", headers=self.headers).json()["results"]
        unit_id = next(item["source_unit_id"] for item in own_search if item["asset_id"] == pdf["id"])
        hidden = self.client.get(
            f"/api/track-d/source-units/{unit_id}", headers={"X-Auth-Token": "other-token"}
        )
        self.assertEqual(hidden.status_code, 404)
        bad = self.client.post(
            "/api/track-d/sources/pdf",
            headers=self.headers,
            files={"file": ("fake.pdf", b"not a pdf", "application/pdf")},
        )
        self.assertEqual(bad.status_code, 422)

    def test_real_local_ocr_extracts_text_without_inventing_description(self):
        result = LocalOCRAnalyzer(timeout_seconds=10).analyze(
            image_bytes=_text_image("ENERGY CONSERVATION"),
            mime_type="image/png",
            source_locator={"kind": "pdf_figure", "page_number": 1},
        )
        self.assertEqual(result.status, "complete", result.error_code)
        self.assertIn("ENERGY", (result.ocr_text or "").upper())
        self.assertIn("CONSERVATION", (result.ocr_text or "").upper())
        self.assertIsNone(result.description)


if __name__ == "__main__":
    unittest.main()
