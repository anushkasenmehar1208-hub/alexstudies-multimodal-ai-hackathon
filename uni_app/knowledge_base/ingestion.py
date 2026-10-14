"""Format-specific extraction into one location-preserving source structure."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from typing import Callable, Iterable, Optional, Sequence
import zipfile

from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy import Engine, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker

from uni_app import youtube_utils

from .models import SourceAsset, SourceUnit, SourceVisual
from .visual_analysis import DeferredVisualAnalyzer, VisualAnalysisResult, VisualAnalyzer


MAX_SOURCE_BYTES = 25 * 1024 * 1024
MAX_PPTX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_OCR_VISUALS_PER_ASSET = 24


class SourceIngestionError(RuntimeError):
    """Base error safe for ingestion callers to classify."""


class InvalidSourceError(SourceIngestionError):
    """The supplied URL or file is not a supported, readable source."""


class DuplicateSourceError(SourceIngestionError):
    """This owner has already ingested the same logical source."""


@dataclass(frozen=True, slots=True)
class _ExtractedUnit:
    unit_index: int
    text: str
    locator: dict[str, object]
    status: str = "complete"


@dataclass(frozen=True, slots=True)
class _ExtractedVisual:
    unit_index: int
    image_index: int
    image_bytes: bytes
    mime_type: str
    width: Optional[int]
    height: Optional[int]
    locator: dict[str, object]


@dataclass(frozen=True, slots=True)
class _ExtractionBundle:
    units: tuple[_ExtractedUnit, ...]
    visuals: tuple[_ExtractedVisual, ...]
    metadata: dict[str, object]
    status: str = "complete"
    error_code: Optional[str] = None


def _digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def _unit_digest(source_type: str, text: str, locator: dict[str, object]) -> str:
    payload = json.dumps(
        {"source_type": source_type, "text": text, "locator": locator},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return _digest(payload)


def _require_owner(owner_id: int) -> int:
    try:
        value = int(owner_id)
    except (TypeError, ValueError) as exc:
        raise InvalidSourceError("owner_id must be a positive integer") from exc
    if value <= 0:
        raise InvalidSourceError("owner_id must be a positive integer")
    return value


def _require_file(filename: str, data: bytes, expected_suffix: str) -> tuple[str, bytes]:
    clean_name = Path(filename or "").name
    if not clean_name or Path(clean_name).suffix.lower() != expected_suffix:
        raise InvalidSourceError(f"expected a {expected_suffix} file")
    if not data:
        raise InvalidSourceError("source file is empty")
    if len(data) > MAX_SOURCE_BYTES:
        raise InvalidSourceError("source file exceeds the 25 MB Phase 1 limit")
    return clean_name, data


def _pdf_mime(extension: str) -> str:
    return {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "jpx": "image/jp2",
        "jp2": "image/jp2",
        "webp": "image/webp",
        "tiff": "image/tiff",
        "bmp": "image/bmp",
    }.get(extension.lower(), f"image/{extension.lower() or 'unknown'}")


def extract_pdf(filename: str, data: bytes) -> _ExtractionBundle:
    clean_name, data = _require_file(filename, data, ".pdf")
    if not data.startswith(b"%PDF-"):
        raise InvalidSourceError("invalid PDF signature")
    try:
        document = PdfReader(BytesIO(data), strict=False)
    except Exception as exc:
        raise InvalidSourceError("PDF could not be opened") from exc
    if document.is_encrypted:
        raise InvalidSourceError("password-protected PDFs are not supported in Phase 1")
    if len(document.pages) <= 0:
        raise InvalidSourceError("PDF contains no pages")

    units: list[_ExtractedUnit] = []
    visuals: list[_ExtractedVisual] = []
    partial = False
    for page_index, page in enumerate(document.pages):
        page_number = page_index + 1
        media_box = page.mediabox
        locator: dict[str, object] = {
            "kind": "pdf_page",
            "page_number": page_number,
            "width_points": round(float(media_box.width), 3),
            "height_points": round(float(media_box.height), 3),
        }
        text = (page.extract_text() or "").strip()
        units.append(_ExtractedUnit(page_index, text, locator))

        try:
            image_keys = list(page.images.keys())
        except Exception:
            image_keys = []
            partial = True
        for image_index, image_key in enumerate(image_keys):
            try:
                image_file = page.images[image_key]
                image_bytes = bytes(image_file.data or b"")
                if not image_bytes:
                    partial = True
                    continue
                name = str(image_file.name or f"image-{image_index}")
                extension = Path(name).suffix.lstrip(".")
                try:
                    pil_image = image_file.image
                    width, height = pil_image.size if pil_image is not None else (None, None)
                except Exception:
                    width, height = None, None
                    partial = True
                visuals.append(
                    _ExtractedVisual(
                        unit_index=page_index,
                        image_index=image_index,
                        image_bytes=image_bytes,
                        mime_type=_pdf_mime(extension),
                        width=int(width) if width is not None else None,
                        height=int(height) if height is not None else None,
                        locator={
                            "kind": "pdf_figure",
                            "page_number": page_number,
                            "image_index": image_index,
                            "object_name": name,
                        },
                    )
                )
            except Exception:
                partial = True

    return _ExtractionBundle(
        units=tuple(units),
        visuals=tuple(visuals),
        metadata={
            "filename": clean_name,
            "page_count": len(units),
            "visual_count": len(visuals),
        },
        status="partial" if partial else "complete",
        error_code="visual_extraction_partial" if partial else None,
    )


def _iter_shapes(shapes: Iterable[object]) -> Iterable[object]:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def extract_pptx(filename: str, data: bytes) -> _ExtractionBundle:
    clean_name, data = _require_file(filename, data, ".pptx")
    archive = BytesIO(data)
    if not zipfile.is_zipfile(archive):
        raise InvalidSourceError("invalid PPTX container")
    archive.seek(0)
    with zipfile.ZipFile(archive) as package:
        if sum(member.file_size for member in package.infolist()) > MAX_PPTX_UNCOMPRESSED_BYTES:
            raise InvalidSourceError("PowerPoint expands beyond the 150 MB safety limit")
        names = set(package.namelist())
        if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
            raise InvalidSourceError("file is not a PowerPoint presentation")
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise InvalidSourceError("PowerPoint file could not be opened") from exc
    if not presentation.slides:
        raise InvalidSourceError("PowerPoint presentation contains no slides")

    units: list[_ExtractedUnit] = []
    visuals: list[_ExtractedVisual] = []
    partial = False
    for slide_index, slide in enumerate(presentation.slides):
        slide_number = slide_index + 1
        text_parts: list[str] = []
        table_parts: list[str] = []
        image_index = 0
        for shape in _iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    text_parts.append(text)
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    table_parts.append("\n".join(rows))
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    width, height = image.size
                    locator = {
                        "kind": "pptx_image",
                        "slide_number": slide_number,
                        "image_index": image_index,
                        "bbox_emu": {
                            "left": int(shape.left),
                            "top": int(shape.top),
                            "width": int(shape.width),
                            "height": int(shape.height),
                        },
                    }
                    visuals.append(
                        _ExtractedVisual(
                            unit_index=slide_index,
                            image_index=image_index,
                            image_bytes=bytes(image.blob),
                            mime_type=str(image.content_type),
                            width=int(width),
                            height=int(height),
                            locator=locator,
                        )
                    )
                    image_index += 1
                except Exception:
                    partial = True

        notes = ""
        try:
            if getattr(slide, "has_notes_slide", False):
                notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                notes = ((notes_frame.text if notes_frame is not None else "") or "").strip()
        except Exception:
            partial = True

        sections: list[str] = []
        if text_parts:
            sections.append("\n".join(text_parts))
        if table_parts:
            sections.append("[Tables]\n" + "\n\n".join(table_parts))
        if notes:
            sections.append("[Speaker notes]\n" + notes)
        locator = {
            "kind": "pptx_slide",
            "slide_number": slide_number,
            "width_emu": int(presentation.slide_width),
            "height_emu": int(presentation.slide_height),
        }
        units.append(_ExtractedUnit(slide_index, "\n\n".join(sections).strip(), locator))

    return _ExtractionBundle(
        units=tuple(units),
        visuals=tuple(visuals),
        metadata={
            "filename": clean_name,
            "slide_count": len(units),
            "visual_count": len(visuals),
        },
        status="partial" if partial else "complete",
        error_code="visual_or_notes_extraction_partial" if partial else None,
    )


TranscriptFetcher = Callable[
    [str],
    tuple[list[youtube_utils.TranscriptSegment], str],
]


class KnowledgeBaseIngestor:
    """Persist source assets, units and visuals with owner-scoped access."""

    def __init__(self, engine: Engine, *, visual_analyzer: Optional[VisualAnalyzer] = None):
        self._sessions = sessionmaker(bind=engine, expire_on_commit=False, class_=Session)
        self._visual_analyzer = visual_analyzer or DeferredVisualAnalyzer()

    def ingest_youtube(
        self,
        *,
        owner_id: int,
        original_url: str,
        title: str = "",
        segments: Optional[Sequence[youtube_utils.TranscriptSegment]] = None,
        fetcher: TranscriptFetcher = youtube_utils.fetch_transcript_segments,
    ) -> SourceAsset:
        owner_id = _require_owner(owner_id)
        video_id = youtube_utils.extract_video_id(original_url)
        if not video_id:
            raise InvalidSourceError("invalid YouTube URL")
        if segments is None:
            fetched, reason = fetcher(video_id)
            if not fetched:
                raise SourceIngestionError(f"transcript unavailable: {reason or 'unknown'}")
            segments = fetched

        units: list[_ExtractedUnit] = []
        for index, segment in enumerate(segments):
            if not isinstance(segment, youtube_utils.TranscriptSegment):
                raise InvalidSourceError("transcript segments must include text, start and end timestamps")
            if not segment.text.strip() or segment.start_seconds < 0 or segment.end_seconds < segment.start_seconds:
                raise InvalidSourceError("transcript contains an invalid timestamped segment")
            locator = {
                "kind": "youtube_timestamp",
                "video_id": video_id,
                "original_url": original_url.strip(),
                "start_seconds": segment.start_seconds,
                "end_seconds": segment.end_seconds,
            }
            units.append(_ExtractedUnit(index, segment.text.strip(), locator))
        if not units:
            raise InvalidSourceError("transcript contains no usable segments")

        bundle = _ExtractionBundle(
            units=tuple(units),
            visuals=(),
            metadata={"video_id": video_id, "segment_count": len(units)},
        )
        identity_hash = _digest(f"youtube:{video_id}".encode("utf-8"))
        return self._persist(
            owner_id=owner_id,
            source_type="youtube",
            title=(title or video_id).strip()[:255],
            original_name=None,
            original_url=original_url.strip(),
            external_id=video_id,
            content_hash=identity_hash,
            bundle=bundle,
        )

    def ingest_pdf(self, *, owner_id: int, filename: str, data: bytes, title: str = "") -> SourceAsset:
        owner_id = _require_owner(owner_id)
        bundle = extract_pdf(filename, data)
        return self._persist(
            owner_id=owner_id,
            source_type="pdf",
            title=(title or Path(filename).stem).strip()[:255],
            original_name=Path(filename).name,
            original_url=None,
            external_id=None,
            content_hash=_digest(data),
            bundle=bundle,
        )

    def ingest_pptx(self, *, owner_id: int, filename: str, data: bytes, title: str = "") -> SourceAsset:
        owner_id = _require_owner(owner_id)
        bundle = extract_pptx(filename, data)
        return self._persist(
            owner_id=owner_id,
            source_type="pptx",
            title=(title or Path(filename).stem).strip()[:255],
            original_name=Path(filename).name,
            original_url=None,
            external_id=None,
            content_hash=_digest(data),
            bundle=bundle,
        )

    def _persist(
        self,
        *,
        owner_id: int,
        source_type: str,
        title: str,
        original_name: Optional[str],
        original_url: Optional[str],
        external_id: Optional[str],
        content_hash: str,
        bundle: _ExtractionBundle,
    ) -> SourceAsset:
        analyzed_visuals: list[tuple[_ExtractedVisual, VisualAnalysisResult]] = []
        ocr_by_unit: dict[int, list[str]] = {}
        for visual_number, extracted in enumerate(bundle.visuals):
            if (
                visual_number >= MAX_OCR_VISUALS_PER_ASSET
                and getattr(self._visual_analyzer, "name", "") != "deferred"
            ):
                analysis = VisualAnalysisResult(
                    status="unavailable",
                    analyzer_name=getattr(self._visual_analyzer, "name", "unknown"),
                    error_code="ocr_asset_limit_reached",
                )
                analyzed_visuals.append((extracted, analysis))
                continue
            try:
                analysis = self._visual_analyzer.analyze(
                    image_bytes=extracted.image_bytes,
                    mime_type=extracted.mime_type,
                    source_locator=extracted.locator,
                )
            except Exception:
                analysis = VisualAnalysisResult(
                    status="failed",
                    analyzer_name=getattr(self._visual_analyzer, "name", "unknown"),
                    error_code="analyzer_failed",
                )
            analyzed_visuals.append((extracted, analysis))
            if analysis.status == "complete" and analysis.ocr_text:
                ocr_by_unit.setdefault(extracted.unit_index, []).append(analysis.ocr_text.strip())

        try:
            with self._sessions.begin() as session:
                duplicate = session.scalar(
                    select(SourceAsset.id).where(
                        SourceAsset.owner_id == owner_id,
                        SourceAsset.source_type == source_type,
                        SourceAsset.content_hash == content_hash,
                    )
                )
                if duplicate is not None:
                    raise DuplicateSourceError("source already exists for this owner")

                asset = SourceAsset(
                    owner_id=owner_id,
                    source_type=source_type,
                    title=title,
                    original_name=original_name,
                    original_url=original_url,
                    external_id=external_id,
                    content_hash=content_hash,
                    extraction_status=bundle.status,
                    error_code=bundle.error_code,
                    source_metadata=bundle.metadata,
                )
                session.add(asset)
                session.flush()

                persisted_units: dict[int, SourceUnit] = {}
                for extracted in bundle.units:
                    unit_text = extracted.text
                    visual_text = [value for value in ocr_by_unit.get(extracted.unit_index, []) if value]
                    if visual_text:
                        suffix = "\n\n[Extracted visual text]\n" + "\n".join(visual_text)
                        unit_text = (unit_text + suffix).strip()
                    unit = SourceUnit(
                        owner_id=owner_id,
                        asset_id=asset.id,
                        unit_index=extracted.unit_index,
                        extracted_text=unit_text,
                        content_hash=_unit_digest(source_type, unit_text, extracted.locator),
                        source_type=source_type,
                        source_locator=extracted.locator,
                        extraction_status=extracted.status,
                    )
                    session.add(unit)
                    persisted_units[extracted.unit_index] = unit
                session.flush()

                for extracted, analysis in analyzed_visuals:
                    unit = persisted_units.get(extracted.unit_index)
                    if unit is None:
                        raise SourceIngestionError("visual references a missing source unit")
                    session.add(
                        SourceVisual(
                            owner_id=owner_id,
                            asset_id=asset.id,
                            source_unit_id=unit.id,
                            unit_index=extracted.unit_index,
                            image_index=extracted.image_index,
                            content_hash=_digest(extracted.image_bytes),
                            mime_type=extracted.mime_type,
                            image_bytes=extracted.image_bytes,
                            width=extracted.width,
                            height=extracted.height,
                            source_locator=extracted.locator,
                            extraction_status="complete",
                            analysis_status=analysis.status,
                            analyzer_name=analysis.analyzer_name,
                            ocr_text=analysis.ocr_text,
                            visual_description=analysis.description,
                        )
                    )
                session.flush()
            return asset
        except DuplicateSourceError:
            raise
        except IntegrityError as exc:
            raise DuplicateSourceError("source already exists for this owner") from exc

    def get_asset(self, *, owner_id: int, asset_id: int) -> Optional[SourceAsset]:
        owner_id = _require_owner(owner_id)
        with self._sessions() as session:
            return session.scalar(
                select(SourceAsset).where(
                    SourceAsset.id == int(asset_id),
                    SourceAsset.owner_id == owner_id,
                )
            )

    def list_assets(self, *, owner_id: int) -> list[SourceAsset]:
        owner_id = _require_owner(owner_id)
        with self._sessions() as session:
            return list(
                session.scalars(
                    select(SourceAsset)
                    .where(SourceAsset.owner_id == owner_id)
                    .order_by(SourceAsset.id)
                )
            )

    def list_units(self, *, owner_id: int, asset_id: int) -> list[SourceUnit]:
        owner_id = _require_owner(owner_id)
        with self._sessions() as session:
            return list(
                session.scalars(
                    select(SourceUnit)
                    .where(SourceUnit.owner_id == owner_id, SourceUnit.asset_id == int(asset_id))
                    .order_by(SourceUnit.unit_index)
                )
            )

    def get_unit(self, *, owner_id: int, source_unit_id: int) -> Optional[SourceUnit]:
        owner_id = _require_owner(owner_id)
        with self._sessions() as session:
            return session.scalar(
                select(SourceUnit).where(
                    SourceUnit.id == int(source_unit_id),
                    SourceUnit.owner_id == owner_id,
                )
            )

    def list_visuals(self, *, owner_id: int, asset_id: int) -> list[SourceVisual]:
        owner_id = _require_owner(owner_id)
        with self._sessions() as session:
            return list(
                session.scalars(
                    select(SourceVisual)
                    .where(SourceVisual.owner_id == owner_id, SourceVisual.asset_id == int(asset_id))
                    .order_by(SourceVisual.unit_index, SourceVisual.image_index)
                )
            )
