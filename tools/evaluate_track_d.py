#!/usr/bin/env python3
"""Run reproducible, local-only Track D retrieval and assessment evaluation."""

from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import sys

from pptx import Presentation
from sqlalchemy import create_engine, delete
from sqlalchemy.engine import make_url

ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(ROOT))

from uni_app import ai_provider, youtube_utils  # noqa: E402
from uni_app.knowledge_base.assessment import AssessmentService  # noqa: E402
from uni_app.knowledge_base.ingestion import KnowledgeBaseIngestor  # noqa: E402
from uni_app.knowledge_base.models import SourceAsset, TopicMastery  # noqa: E402
from uni_app.knowledge_base.retrieval import SourceRetriever  # noqa: E402
from uni_app.knowledge_base.tutoring import GroundedTutor  # noqa: E402
from uni_app.knowledge_base.visual_analysis import DeferredVisualAnalyzer  # noqa: E402
from tools.pdf_fixture import build_pdf  # noqa: E402

OWNER = 1_904_001


def engine():
    url = make_url(os.environ.get("DATABASE_URL", ""))
    if (url.host, url.port, url.database) not in {
        ("127.0.0.1", 55432, "alexstudies_hackathon"),
        ("localhost", 55432, "alexstudies_hackathon"),
    }:
        raise RuntimeError("Evaluation refuses any database except local alexstudies_hackathon:55432")
    return create_engine(url.set(drivername="postgresql+psycopg"), pool_pre_ping=True)


def pdf_bytes():
    return build_pdf([(612, 792, "Chloroplast membranes perform photophosphorylation. Photosystems capture light energy. ATP stores chemical energy. Water supplies electrons. Pure water boils at 100 degrees Celsius.", None)])


def pptx_bytes():
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = "Impulse and momentum"
    slide.shapes.add_textbox(914400, 1371600, 6400800, 1371600).text = "Impulse changes momentum during a collision. Force acts over time. Momentum is conserved in an isolated collision."
    output = BytesIO(); deck.save(output); return output.getvalue()


def main():
    db = engine(); ingestor = KnowledgeBaseIngestor(db, visual_analyzer=DeferredVisualAnalyzer())
    with db.begin() as connection:
        connection.execute(delete(TopicMastery).where(TopicMastery.owner_id == OWNER))
        connection.execute(delete(SourceAsset).where(SourceAsset.owner_id == OWNER))
    try:
        ingestor.ingest_pdf(owner_id=OWNER, filename="photosynthesis.pdf", data=pdf_bytes())
        ingestor.ingest_pptx(owner_id=OWNER, filename="mechanics.pptx", data=pptx_bytes())
        ingestor.ingest_youtube(
            owner_id=OWNER,
            original_url="https://www.youtube.com/watch?v=abcdefghijk",
            title="Protein synthesis",
            segments=[youtube_utils.TranscriptSegment("A ribosome reads each codon during translation. Transfer RNA carries amino acids. Proteins form from linked amino acids.", 42.0, 52.0)],
        )
        fixture = json.loads((ROOT / "tests/fixtures/track_d_evaluation.json").read_text())
        retriever = SourceRetriever(db)
        hits = []
        reciprocal_ranks = []
        precisions = []
        recalls = []
        for case in fixture["retrieval_cases"]:
            results = retriever.search(owner_id=OWNER, query=case["query"], limit=3)
            expected = [index for index, item in enumerate(results, 1) if item.source_type == case["expected_source_type"]]
            hits.append(bool(expected)); reciprocal_ranks.append(1 / expected[0] if expected else 0.0)
            precisions.append(sum(item.source_type == case["expected_source_type"] for item in results) / max(len(results), 1))
            recalls.append(1.0 if expected else 0.0)

        provider_called = False
        def forbidden_provider(*_args, **_kwargs):
            nonlocal provider_called; provider_called = True; raise AssertionError("provider called for unsupported query")
        config = ai_provider.ProviderConfig(True, "groq", "", "https://api.groq.com/openai/v1", "openai/gpt-oss-20b", "whisper-large-v3-turbo")
        refusal = GroundedTutor(retriever, config=config, completion=forbidden_provider).answer(owner_id=OWNER, question=fixture["unsupported_question"])

        assessments = AssessmentService(db)
        generated = assessments.generate(owner_id=OWNER, limit=15, diagnostic=True)
        repeated = assessments.generate(owner_id=OWNER, limit=15, diagnostic=True)
        citation_validity = sum(item.citation.get("source_unit_id") == item.source_unit_id for item in generated.questions) / len(generated.questions)
        unique_prompts = len({item.prompt for item in generated.questions})
        first = generated.questions[0]
        assessments.answer(owner_id=OWNER, question_id=first.id, answer=first.correct_answer)
        assessments.answer(owner_id=OWNER, question_id=first.id, answer="deliberately incorrect")
        mastery = next(row for row in assessments.mastery(owner_id=OWNER) if row["topic_tag"] == first.topic_tag)

        report = {
            "source_fixture_count": 3,
            "fixture_cases": len(hits),
            "retrieval_hit_rate": round(sum(hits) / len(hits), 4),
            "mean_reciprocal_rank": round(sum(reciprocal_ranks) / len(reciprocal_ranks), 4),
            "context_precision_at_3": round(sum(precisions) / len(precisions), 4),
            "context_recall_at_3": round(sum(recalls) / len(recalls), 4),
            "citation_validity": round(citation_validity, 4),
            "assessment_question_count": len(generated.questions),
            "unsupported_refusal": not refusal.supported and not refusal.citations and not provider_called,
            "question_repetition_rate": round(1 - unique_prompts / len(generated.questions), 4),
            "duplicate_generation_created_count": repeated.created_count,
            "simulated_mastery_after_correct_then_wrong": mastery["mastery_score"],
            "question_types": sorted({item.question_type for item in generated.questions}),
            "faithfulness": "not_run_no_development_provider_key",
            "answer_relevance": "not_run_no_development_provider_key"
        }
        output = ROOT / "reports/track_d_evaluation.json"; output.parent.mkdir(exist_ok=True); output.write_text(json.dumps(report, indent=2) + "\n")
        print(json.dumps(report, indent=2))
    finally:
        with db.begin() as connection:
            connection.execute(delete(TopicMastery).where(TopicMastery.owner_id == OWNER))
            connection.execute(delete(SourceAsset).where(SourceAsset.owner_id == OWNER))
        db.dispose()


if __name__ == "__main__":
    main()
