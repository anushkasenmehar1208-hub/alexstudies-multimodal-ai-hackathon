"""Create small deterministic Track D demo files outside the tracked tree."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation

from pdf_fixture import build_pdf


def create_fixtures(output_directory: Path) -> tuple[Path, Path]:
    output_directory.mkdir(parents=True, exist_ok=True)

    pdf_path = output_directory / "photosynthesis-textbook.pdf"
    pdf_path.write_bytes(
        build_pdf(
            [
                (
                    612,
                    792,
                    "Photosynthesis study text\n"
                    "Chloroplast membranes perform the light-dependent reactions.\n"
                    "These reactions use light energy to produce ATP and NADPH.",
                    None,
                ),
                (
                    612,
                    792,
                    "The Calvin cycle uses ATP and NADPH to help build sugars from carbon dioxide.",
                    None,
                ),
            ]
        )
    )

    pptx_path = output_directory / "momentum-slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Impulse and momentum"
    slide.placeholders[1].text = (
        "Impulse equals force multiplied by the time interval.\n"
        "Impulse changes an object's momentum."
    )
    slide.notes_slide.notes_text_frame.text = "Use the collision example during the demonstration."

    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Momentum relationship"
    second.placeholders[1].text = "Momentum equals mass multiplied by velocity."
    presentation.save(pptx_path)
    return pdf_path, pptx_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(".local-data/demo-fixtures"),
        help="Output directory (default: ignored local data directory)",
    )
    args = parser.parse_args()
    pdf_path, pptx_path = create_fixtures(args.output.resolve())
    print(f"Created {pdf_path}")
    print(f"Created {pptx_path}")


if __name__ == "__main__":
    main()
